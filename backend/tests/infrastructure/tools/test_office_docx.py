"""`.docx` 真解析：按文档真实顺序切块（段落与表格交错走）。

fixture 是**真 Word 文件**（python-docx 现造的字节，不是替身）——解析器真读它。
"""

import io

import pytest
from docx import Document

from pmstudio.application.errors import UnsupportedMaterial
from pmstudio.infrastructure.tools.extract import MAX_MATERIAL_CHARS, LocalMaterialParser
from pmstudio.infrastructure.tools.office import docx_chunks

parser = LocalMaterialParser()


def docx_bytes(*, heading: str = "导出权限", body: str = "只有管理员能导出。",
               with_table: bool = True) -> bytes:
    """一份最小真 docx：标题 + 正文 +（可选）一行表头的表格 + 二级标题 + 正文。"""

    document = Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(body)
    if with_table:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "角色"
        table.cell(0, 1).text = "能否导出"
        table.cell(1, 0).text = "管理员"
        table.cell(1, 1).text = "能"
    document.add_heading("边界", level=2)
    document.add_paragraph("普通成员看不到按钮。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_keeps_paragraph_and_table_order():
    """段落与表格**交错**：标题起新块、表格独占一块、锚点是"第 N 段"与"表格 N"。"""

    chunks, truncated = docx_chunks(docx_bytes())

    assert truncated is False
    assert [chunk.ordinal for chunk in chunks] == [0, 1, 2]
    first = chunks[0]
    assert first.text.splitlines()[0] == "导出权限"
    assert "只有管理员能导出。" in first.text
    assert first.heading_path == ("导出权限",)
    assert first.anchor == "第 1–2 段"
    table = chunks[1]
    assert table.text == "角色 | 能否导出\n管理员 | 能"
    assert table.anchor == "表格 1（表头：角色、能否导出）"
    assert table.heading_path == ("导出权限",)
    last = chunks[2]
    assert last.heading_path == ("导出权限", "边界")
    assert "普通成员看不到按钮。" in last.text


def test_docx_becomes_an_upload_material_with_its_own_identity():
    parsed = parser.parse(
        filename="需求.docx",
        media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        data=docx_bytes(),
    )

    assert parsed.origin == "upload"
    assert parsed.source_ref == "attachment:需求.docx"
    assert parsed.parser_id == "python-docx"
    assert parsed.parser_version.startswith("python-docx ")
    assert parsed.truncated is False
    assert parsed.warnings == ()


def test_an_empty_docx_says_it_has_no_text():
    """空 Word ≠ 空材料：照实说"没有可抽取的文字"。"""

    buffer = io.BytesIO()
    Document().save(buffer)

    with pytest.raises(UnsupportedMaterial) as error:
        parser.parse(
            filename="空.docx",
            media_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            data=buffer.getvalue(),
        )

    assert "没有可抽取的文字" in str(error.value)


def test_a_broken_docx_is_a_clear_refusal_not_a_500():
    """坏文件（不是 zip / 半截）：报"读不了"，不许把 BadZipFile 漏到 HTTP 层。"""

    with pytest.raises(UnsupportedMaterial) as error:
        parser.parse(
            filename="broken.docx",
            media_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            data=b"not a docx",
        )

    assert error.value.code == "unsupported_material"
    assert "读不了" in str(error.value)


def test_a_very_long_docx_is_marked_truncated():
    document = Document()
    for _ in range(400):
        document.add_paragraph("字" * 1000)
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = parser.parse(
        filename="大.docx",
        media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        data=buffer.getvalue(),
    )

    assert parsed.truncated is True
    assert any("截断" in warning for warning in parsed.warnings)
    assert sum(len(chunk.text) for chunk in parsed.chunks) <= MAX_MATERIAL_CHARS + 500
