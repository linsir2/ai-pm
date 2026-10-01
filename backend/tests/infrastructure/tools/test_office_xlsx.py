"""`.xlsx` 真解析：一个工作表一块，行按 `a | b` 渲染，锚点写工作表名与行范围。

fixture 是**真 Excel 文件**（openpyxl 现造的字节，不是替身）——解析器真读它。
"""

import io

import openpyxl
import pytest

from pmstudio.application.errors import UnsupportedMaterial
from pmstudio.infrastructure.tools.extract import MAX_MATERIAL_CHARS, LocalMaterialParser
from pmstudio.infrastructure.tools.office import xlsx_chunks

parser = LocalMaterialParser()

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx_bytes(*sheets: tuple[str, list[list[object]]]) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        sheet = workbook.create_sheet(title=title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_each_sheet_becomes_one_chunk_with_its_own_anchor():
    chunks, truncated, warnings = xlsx_chunks(
        xlsx_bytes(
            ("指标", [["指标", "目标"], ["导出耗时", "< 3s"]]),
            ("明细", [["角色", "能否导出"], ["管理员", "能"]]),
        )
    )

    assert truncated is False
    assert warnings == ()
    assert [chunk.text for chunk in chunks] == [
        "指标 | 目标\n导出耗时 | < 3s",
        "角色 | 能否导出\n管理员 | 能",
    ]
    assert [chunk.anchor for chunk in chunks] == [
        "工作表「指标」第 1–2 行",
        "工作表「明细」第 1–2 行",
    ]
    assert [chunk.ordinal for chunk in chunks] == [0, 1]
    assert chunks[0].heading_path == ("指标",)


def test_empty_sheets_are_skipped_not_stored_as_empty_chunks():
    chunks, _, _ = xlsx_chunks(
        xlsx_bytes(("空表", []), ("有货", [["甲", "乙"]]))
    )

    assert [chunk.text for chunk in chunks] == ["甲 | 乙"]
    assert chunks[0].anchor == "工作表「有货」第 1 行"


def test_a_workbook_without_any_cell_says_it_has_no_text():
    with pytest.raises(UnsupportedMaterial) as error:
        parser.parse(filename="空.xlsx", media_type=XLSX_MEDIA_TYPE, data=xlsx_bytes(("空", [])))

    assert "没有可抽取的文字" in str(error.value)


def test_an_unreadable_xlsx_is_refused_not_a_500():
    with pytest.raises(UnsupportedMaterial) as error:
        parser.parse(filename="broken.xlsx", media_type=XLSX_MEDIA_TYPE, data=b"not a zip")

    assert error.value.code == "unsupported_material"
    assert "读不了" in str(error.value)


def test_a_very_long_sheet_is_marked_truncated():
    rows = [["字" * 1000] for _ in range(400)]

    parsed = parser.parse(
        filename="大.xlsx", media_type=XLSX_MEDIA_TYPE, data=xlsx_bytes(("长表", rows))
    )

    assert parsed.truncated is True
    assert any("截断" in warning for warning in parsed.warnings)
    assert sum(len(chunk.text) for chunk in parsed.chunks) <= MAX_MATERIAL_CHARS + 100


def test_xlsx_becomes_an_upload_material_with_its_own_identity():
    parsed = parser.parse(
        filename="排期.xlsx",
        media_type=XLSX_MEDIA_TYPE,
        data=xlsx_bytes(("排期", [["任务", "截止"]])),
    )

    assert parsed.origin == "upload"
    assert parsed.source_ref == "attachment:排期.xlsx"
    assert parsed.parser_id == "openpyxl"
    assert parsed.parser_version.startswith("openpyxl ")


def test_formula_cells_without_a_cached_value_are_reported():
    """公式没算过 → 那格是空的：**照实说一句**，别让人以为这格本来就没内容。"""

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "算账"
    sheet.append(["合计"])
    sheet["A2"] = "=SUM(1, 2)"
    buffer = io.BytesIO()
    workbook.save(buffer)

    chunks, _, warnings = xlsx_chunks(buffer.getvalue())

    assert [chunk.text for chunk in chunks] == ["合计"]
    assert any("公式" in warning for warning in warnings)
